"""Teste ponta-a-ponta LOCAL: navegador real + handler real de api/analyze.js + IA simulada + CSP de vercel.json aplicada.
Uso (na raiz do projecto): python3 scripts/e2e.py   Requer: node, playwright, reportlab, python-docx, pillow.
Não testa: fornecedor de IA real, Redis real, Vercel, pagamentos/WhatsApp."""
import subprocess,time,json,os,sys,urllib.request,urllib.error,tempfile,pathlib
from playwright.sync_api import sync_playwright
import socket,atexit
def free_port():
    s=socket.socket();s.bind(('127.0.0.1',0));p=s.getsockname()[1];s.close();return p
ROOT=pathlib.Path(__file__).resolve().parents[1];PORT=free_port();BASE=f'http://localhost:{PORT}'  # porta livre: nunca reutiliza um servidor antigo (os contadores por código são em memória)
LOG=pathlib.Path(tempfile.gettempdir())/'applypro-dev.log'
ENV=dict(os.environ,STUB_AI='1',PREMIUM_CODES='maria=MARIA-CODE-1234,joao=JOAO-CODE-5678,ana=ANA-CODE-4242,LEGACYCODE12345',PREMIUM_DAILY_LIMIT='4',AI_API_KEY='stub-key-not-real')
srv=subprocess.Popen(['node',str(ROOT/'scripts/dev-server.mjs'),str(PORT)],cwd=ROOT,env=ENV,stdout=open(LOG,'w'),stderr=subprocess.STDOUT)
atexit.register(lambda:(srv.terminate(),srv.wait(5)))  # termina sempre o servidor, mesmo se o teste falhar a meio
# ---- ficheiros de teste reais
FX=pathlib.Path(tempfile.mkdtemp())
CVTXT=['Joao Silva (ficticio) - Licenciado em Letras','Professor de ingles durante 4 anos numa escola secundaria.','Traducao tecnica EN-PT. Conhecimentos de Word e Excel. MARCADOR-E2E']
from reportlab.pdfgen import canvas
c=canvas.Canvas(str(FX/'cv.pdf'))
for i,l in enumerate(CVTXT):c.drawString(72,780-20*i,l)
c.save()
c=canvas.Canvas(str(FX/'long.pdf'))
for p in range(16):c.drawString(72,780,f'Pagina {p+1} com texto suficiente para o teste de limite de paginas do PDF.');c.showPage()
c.save()
import docx;d=docx.Document()
for l in CVTXT:d.add_paragraph(l)
d.save(str(FX/'cv.docx'))
from PIL import Image,ImageDraw
im=Image.new('RGB',(600,300),'white');ImageDraw.Draw(im).text((20,20),'texto em imagem',fill='black');im.save(FX/'scan.pdf')
(FX/'fake.pdf').write_text('isto nao e um pdf');(FX/'cv.txt').write_text('x'*80)
# ---- fotografias de teste de um "CV em papel" (rodada, com ruído e luz desigual), geradas aqui
def make_scan_fixtures(d):
    import random,numpy as np
    from PIL import Image,ImageDraw,ImageFont,ImageFilter
    fp=[(a,b) for a in ['/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf','/usr/share/fonts/dejavu/DejaVuSans.ttf','/Library/Fonts/Arial.ttf','C:/Windows/Fonts/arial.ttf'] for b in [a.replace('Sans.ttf','Sans-Bold.ttf')]]
    fp=[x for x in fp if pathlib.Path(x[0]).exists()]
    if not fp:return False
    F,FB=fp[0];random.seed(3);np.random.seed(3)
    def page(lines,w=1654,h=2339):
        im=Image.new('RGB',(w,h),(250,250,247));dr=ImageDraw.Draw(im);y=110
        for txt,bold,size in lines:
            dr.text((120,y),txt,fill=(25,25,30),font=ImageFont.truetype(FB if bold else F,size));y+=int(size*1.55)
        return im
    def photo(im,ang,out):
        a=np.array(im).astype(np.float32);h,w,_=a.shape;a=a*(np.linspace(0.82,1.05,w)[None,:,None]*np.linspace(1.0,0.9,h)[:,None,None])+np.random.normal(0,5,a.shape)
        r=Image.fromarray(np.clip(a,0,255).astype('uint8')).rotate(ang,expand=True,fillcolor=(70,62,55),resample=Image.BICUBIC).filter(ImageFilter.GaussianBlur(0.7))
        r.resize((int(r.width*0.9),int(r.height*0.9))).save(out,quality=82)
    L1=[("Joao Silva",1,54),("Professor de inglês e tradutor técnico",0,32),("",0,20),("RESUMO",1,30),("Licenciado em Letras com quatro anos de experiência como professor",0,26),("de inglês e em tradução técnica entre inglês e português.",0,26),("",0,20),("EXPERIÊNCIA PROFISSIONAL",1,30),("Professor de inglês - Escola Secundária de Maputo (2020 a 2024)",1,26),("Leccionou inglês a turmas do ensino secundário.",0,26),("",0,20),("FORMAÇÃO",1,30),("Licenciatura em Letras - Universidade Eduardo Mondlane",0,26)]
    L2=[("CERTIFICADOS E INFORMAÇÃO ADICIONAL",1,30),("Curso de tradução técnica - Centro de Línguas (2022)",0,26),("Línguas: Português (nativo), Inglês (avançado).",0,26)]
    photo(page(L1),1.6,d/'cv-photo.jpg');photo(page(L2),-1.2,d/'cv-photo-p2.jpg')
    Image.open(d/'cv-photo.jpg').rotate(90,expand=True).save(d/'cv-photo-rot90.jpg',quality=82)
    Image.open(d/'cv-photo.jpg').save(d/'cv-photo.png');Image.open(d/'cv-photo.jpg').save(d/'cv-photo.webp',quality=80)
    Image.new('RGB',(1200,1600),(235,235,230)).save(d/'blank.jpg',quality=85)
    (d/'fake.jpg').write_bytes(b'%PDF-1.4 isto nao e uma imagem');(d/'big.jpg').write_bytes(b'\xff\xd8\xff'+b'0'*(16*1024*1024))
    return True
HAVE_SCAN=make_scan_fixtures(FX)
def scan_done(pg,timeout=120000):
    until(pg,"!S.scan.busy&&(S.scan.msg||S.scan.err)",timeout)
R=[]
def pdftext(path):
    return subprocess.run(['pdftotext','-layout',str(path),'-'],capture_output=True,text=True).stdout
def docxtext(path):
    import docx as _d;return '\n'.join(p.text for p in _d.Document(str(path)).paragraphs)
def download(pg,selector):
    with pg.expect_download() as d:pg.click(selector)
    dl=d.value;p=FX/('dl-'+dl.suggested_filename);dl.save_as(str(p));return dl.suggested_filename,p
def until(pg,expr,timeout=15000):
    # Polling via CDP (pg.evaluate). NÃO usar pg.wait_for_function: avalia texto com eval() na página, que a CSP bloqueia (artefacto do teste, não da app).
    t0=time.time()
    while (time.time()-t0)*1000<timeout:
        if pg.evaluate('()=>!!('+expr+')'):return
        time.sleep(.1)
    raise TimeoutError(expr)
def check(name,ok,info=''):R.append((name,bool(ok)));print(('PASS' if ok else 'FAIL'),name,info)
def post(body,ip):
    r=urllib.request.Request(BASE+'/api/analyze',json.dumps(body).encode(),{'content-type':'application/json','x-real-ip':ip})
    try:
        with urllib.request.urlopen(r) as x:return x.status,json.loads(x.read())
    except urllib.error.HTTPError as e:return e.code,json.loads(e.read())
for _ in range(50):
    try:urllib.request.urlopen(BASE+'/');break
    except Exception:time.sleep(.2)
CV='Joao Silva (ficticio). Licenciado em Letras. Professor de ingles durante 4 anos. Traducao tecnica EN-PT.';VAC='Assistente de Projectos. Requisitos: licenciatura, ingles fluente, gestao de projectos, Excel.'
with sync_playwright() as p:
    b=p.chromium.launch()
    # nota: o 402 de um código inválido aparece na consola do navegador; é esperado nos testes e por isso é ignorado
    def ctx(ip,w=1280,h=900):
        cx=b.new_context(viewport={'width':w,'height':h},extra_http_headers={'x-real-ip':ip});pg=cx.new_page();pg.errs=[]
        pg.add_init_script("window.__csp=[];document.addEventListener('securitypolicyviolation',e=>window.__csp.push(e.violatedDirective+' '+e.blockedURI))")
        pg.on('console',lambda m:pg.errs.append(m.text) if m.type=='error' and 'status of 402' not in m.text else None);pg.on('pageerror',lambda e:pg.errs.append(str(e)));return pg
    def clean(pg,name):
        v=pg.evaluate('window.__csp');check(name+': sem violações CSP nem erros de consola',not v and not pg.errs,f'csp={v} errs={pg.errs}')
    # 1) cabeçalhos + landing PT
    r=urllib.request.urlopen(BASE+'/');csp=r.headers.get('Content-Security-Policy','')
    import re as _re
    check("CSP presente: script-src só 'self' + 'wasm-unsafe-eval' (OCR); sem 'unsafe-inline' nem 'unsafe-eval'",_re.search(r"script-src 'self' 'wasm-unsafe-eval'(;|$)",csp) is not None and not _re.search(r"'unsafe-(inline|eval)'",csp))
    pg=ctx('10.1.0.1');pg.goto(BASE+'/');pg.wait_for_selector('.hero')
    check('landing PT: h1 e title',pg.inner_text('h1').startswith('O teu CV') and pg.title().startswith('ApplyPro'))
    check('fonte Plus Jakarta Sans carregada',pg.evaluate("document.fonts.check('800 16px \"Plus Jakarta Sans\"')"))
    check('anel de pontuação recebe --p por CSSOM',pg.evaluate("getComputedStyle(document.querySelector('.ring')).getPropertyValue('--p').trim()")=='78')
    # 2) idioma
    pg.click('[data-lang=en]');until(pg,"document.querySelector('h1').innerText.startsWith('Is your CV')")
    check('EN: h1, html lang, title, footer traduzidos',pg.evaluate("document.documentElement.lang")=='en' and pg.title().startswith('ApplyPro | AI-powered') and 'Privacy' in pg.inner_text('footer'))
    pg.wait_for_timeout(300)
    check('EN: imagem hero EN carrega',pg.evaluate("(()=>{const i=document.querySelector('.heroimg');return i.currentSrc.includes('-en-')&&i.naturalWidth>0})()"))
    pg.reload();pg.wait_for_selector('.hero');check('EN persiste após recarregar (localStorage)',pg.inner_text('h1').startswith('Is your CV'))
    pg.goto(BASE+'/?lang=pt');pg.wait_for_selector('.hero');check('?lang=pt sobrepõe a escolha guardada',pg.inner_text('h1').startswith('O teu CV'))
    clean(pg,'landing')
    # 3) upload real PDF/DOCX sob CSP
    for fn,lab in [('cv.pdf','PDF'),('cv.docx','DOCX')]:
        pg=ctx('10.1.0.2');pg.goto(BASE+'/');pg.click('[data-act=start]');pg.set_input_files('#f',str(FX/fn))
        try:until(pg,"document.querySelector('#cv').value.includes('MARCADOR-E2E')",timeout=20000);ok=True
        except Exception:ok=False
        check(f'upload {lab} real: texto extraído para a caixa do CV',ok,pg.inner_text('.err') if not ok else '');clean(pg,'upload '+lab)
    for fn,frag,lab in [('scan.pdf','digitalizado','PDF só com imagem'),('fake.pdf','não corresponde','conteúdo falso .pdf'),('long.pdf','páginas','PDF > 15 páginas'),('cv.txt','Formato não suportado','extensão .txt')]:
        pg=ctx('10.1.0.3');pg.goto(BASE+'/');pg.click('[data-act=start]');pg.set_input_files('#f',str(FX/fn));until(pg,"!document.querySelector('.err').classList.contains('hide')",timeout=20000)
        check(f'recusa: {lab}',frag in pg.inner_text('.err'),pg.inner_text('.err')[:60])
    pg=ctx('10.1.0.4');pg.goto(BASE+'/');pg.click('[data-act=start]');pg.focus('#f')
    check('input de ficheiro é focável por teclado',pg.evaluate("document.activeElement.id")=='f')
    # 4) fluxo completo PT
    pg=ctx('10.1.1.1');pg.goto(BASE+'/');pg.click('[data-act=start]');pg.fill('#cv',CV);pg.click('[data-act=next1]');pg.fill('#vac',VAC);pg.click('[data-act=next2]');pg.wait_for_selector('.score',timeout=15000)
    check('PT grátis: pontuação 64/100 e só 1 melhoria',pg.inner_text('.score .ring i')=='64/100' and pg.locator('text=Actual:').count()==1)
    pg.click('[data-act=kit]');check('PT: kit bloqueado sem código',pg.locator('.lock').count()==1)
    pg.fill('#code','codigo-errado');pg.click('[data-act=unlock]');pg.wait_for_selector('.lock .err:not(.hide)');check('PT: código inválido => erro',pg.inner_text('.lock .err')=='Código Premium inválido.')
    pg.fill('#code','MARIA-CODE-1234');pg.press('#code','Enter');pg.wait_for_selector('[data-act=genkit]');pg.click('[data-act=genkit]');pg.wait_for_selector('#k1');
    check('PT: código pessoal (Enter) desbloqueia e gera o kit',pg.inner_text('#k1').startswith('Carta [STUB]'))
    pg.click('[data-act=copy]');pg.click('[data-act=iv]');pg.click('[data-act=geniv]');pg.wait_for_selector('details summary')
    check('PT: entrevista com exemplo por defeito',pg.inner_text('details summary').startswith('Pergunta 1'))
    clean(pg,'fluxo PT')
    # 5) fluxo completo EN (kit/entrevista pedidos em inglês)
    pg=ctx('10.1.2.1');pg.goto(BASE+'/?lang=en');pg.click('[data-act=start]');pg.fill('#cv',CV);pg.click('[data-act=next1]');pg.click('[data-act=next2]')
    check('EN: validação traduzida',pg.inner_text('.err')=='Paste the job description to continue.')
    pg.fill('#vac',VAC);pg.click('[data-act=next2]');pg.wait_for_selector('.score',timeout=15000);check('EN: resultados em inglês',pg.locator('text=Strengths').count()==1 and pg.locator('text=Fluent English').count()>=1)
    pg.click('[data-act=kit]');pg.fill('#code','x');pg.click('[data-act=unlock]');pg.wait_for_selector('.lock .err:not(.hide)');check('EN: código inválido em inglês',pg.inner_text('.lock .err')=='Invalid Premium code.')
    pg.fill('#code','JOAO-CODE-5678');pg.click('[data-act=unlock]');pg.wait_for_selector('[data-act=genkit]');pg.click('[data-act=genkit]');pg.wait_for_selector('#k1');pg.click('[data-act=iv]');pg.click('[data-act=geniv]');pg.wait_for_selector('details summary')
    pg.click('details summary');check('EN: entrevista, exemplo por defeito em inglês',pg.locator('text=Not identified in the CV.').count()>=1);clean(pg,'fluxo EN')
    calls=json.loads(urllib.request.urlopen(BASE+'/__calls').read())
    check('o pedido EN chegou ao fornecedor com a instrução de língua; PT não',[c['lang'] for c in calls if c['task']=='kit']==['pt','en'] and calls[0]['lang']=='pt',str([(c['task'],c['lang']) for c in calls]))
    # 6) limite diário por código (4/dia) e códigos pessoais
    ip='10.1.3.1';sts=[post({'task':'analyze','cv':CV,'vacancy':VAC,'code':'LEGACYCODE12345'},ip)[0] for _ in range(5)]
    check('código antigo (sem etiqueta) continua válido; 5.º pedido => 429',sts==[200,200,200,200,429],str(sts))
    s,j=post({'task':'analyze','cv':CV,'vacancy':VAC,'code':'LEGACYCODE12345'},'10.1.3.2');check('limite é por código (não por IP): outro IP também bloqueado',s==429 and j['code']=='CODE_LIMIT')
    s,j=post({'task':'analyze','cv':CV,'vacancy':VAC,'code':'JOAO-CODE-5678'},'10.1.3.3');check('outro código não é afectado',s==200)
    s,j=post({'task':'verify','code':'LEGACYCODE12345'},'10.1.3.4');check('verify não conta para o limite diário',s==200)
    s,j=post({'task':'verify','code':'maria=MARIA-CODE-1234'},'10.1.3.5');check('entrada completa "etiqueta=código" também é aceite (compatibilidade)',s==200)
    s,j=post({'task':'analyze','cv':CV,'vacancy':VAC,'code':'maria'},'10.1.3.6');check('a etiqueta sozinha NÃO é um código',s==402)
    # 6b) Voltar em todos os passos + botão voltar do navegador
    pg=ctx('10.1.6.1');pg.goto(BASE+'/');pg.click('[data-act=start]');pg.fill('#cv',CV);pg.click('[data-act=next1]');pg.fill('#vac',VAC);pg.click('[data-act=next2]');pg.wait_for_selector('.score',timeout=15000)
    pg.click('[data-act=kit]');pg.wait_for_selector('.lock');check('passo 4 tem botão Voltar',pg.locator('[data-act=goto][data-step="3"]').count()==1)
    pg.click('[data-act=goto][data-step="3"]');pg.wait_for_selector('.score');check('Voltar (4→3) mostra a análise sem a repetir',pg.inner_text('.score .ring i')=='64/100')
    pg.click('[data-act=goto][data-step="2"]');pg.wait_for_selector('#vac');check('Voltar (3→2) mantém o texto da vaga',pg.input_value('#vac')==VAC)
    pg.click('[data-act=goto][data-step="1"]');pg.wait_for_selector('#cv');check('Voltar (2→1) mantém o texto do CV',pg.input_value('#cv')==CV)
    pg.click('[data-act=goto][data-step="0"]');pg.wait_for_selector('.hero');check('Voltar (1→início) mostra a landing',pg.locator('#sobre').count()==1)
    pg.click('[data-act=start]');pg.click('[data-act=next1]');pg.wait_for_selector('#vac');pg.go_back();pg.wait_for_selector('#cv')
    check('botão voltar do navegador recua um passo (2→1) em vez de sair da app',pg.url.rstrip('/')==BASE and pg.locator('#cv').count()==1)
    pg.go_forward();pg.wait_for_selector('#vac');check('avançar do navegador volta ao passo 2',pg.locator('#vac').count()==1)
    pg=ctx('10.1.6.7');pg.goto(BASE+'/');pg.click('.hero >> text=Como funciona');pg.click('[data-act=start]');pg.wait_for_selector('#cv');pg.go_back();pg.wait_for_selector('.hero')
    check('voltar do passo 1 para uma entrada de âncora mostra a landing',pg.locator('.hero').count()==1)
    # 6c) Kit em PDF/Word + CV adaptado (editar, rever, descarregar)
    pg=ctx('10.1.6.2');pg.goto(BASE+'/');pg.click('[data-act=start]');pg.fill('#cv',CV);pg.click('[data-act=next1]');pg.fill('#vac',VAC);pg.click('[data-act=next2]');pg.wait_for_selector('.score',timeout=15000)
    pg.click('[data-act=kit]');pg.fill('#code','ANA-CODE-4242');pg.click('[data-act=unlock]');pg.wait_for_selector('[data-act=genkit]');pg.click('[data-act=genkit]');pg.wait_for_selector('#k1')
    check('as 3 caixas do kit têm PDF e Word',pg.locator('[data-act=dl][data-fmt=pdf]').count()==3 and pg.locator('[data-act=dl][data-fmt=docx]').count()==3)
    fn,path=download(pg,'[data-act=dl][data-id=k1][data-fmt=pdf]');txt=pdftext(path)
    from pypdf import PdfReader
    r=PdfReader(str(path));pw,ph=float(r.pages[0].mediabox.width),float(r.pages[0].mediabox.height)
    check('carta em PDF: ficheiro %PDF, A4, texto da carta',open(path,'rb').read(4)==b'%PDF' and abs(pw-595)<2 and abs(ph-842)<2 and 'Carta [STUB]' in txt and fn=='Carta-de-apresentacao.pdf',f'{fn} {pw}x{ph}')
    fn,path=download(pg,'[data-act=dl][data-id=k1][data-fmt=docx]');check('carta em Word: .docx abre no python-docx com o texto',fn.endswith('.docx') and 'Carta [STUB]' in docxtext(path),fn)
    fn,path=download(pg,'[data-act=dl][data-id=k2][data-fmt=pdf]');t2=pdftext(path);check('email em PDF: assunto e corpo',('Assunto: Assunto' in t2) and 'Corpo' in t2)
    check('antes de gerar: cartão do CV com botão "Gerar CV adaptado", sem formulário',pg.locator('[data-act=gencv]').count()==1 and pg.locator('.cvform').count()==0)
    pg.click('[data-act=gencv]');pg.wait_for_selector('.cvform',timeout=15000)
    wt=pg.inner_text('.cvform .warn:nth-of-type(2)') if pg.locator('.cvform .warn').count()>1 else ''
    allw=' '.join(pg.locator('.cvform .warn').all_inner_texts())
    check('verificação: contacto inventado removido, nome fiel mantido',pg.input_value('[data-cv=contact]')=='' and pg.input_value('[data-cv=name]')=='Joao Silva')
    check('verificação: número (120) e nomes (Escola Y, UEM) assinalados ao utilizador','120' in allw and 'Escola Y' in allw and 'UEM' in allw,allw[:160])
    check('descarregar está desactivado até confirmar a revisão',pg.is_disabled('[data-act=cvdl][data-fmt=pdf]') and pg.is_disabled('[data-act=cvdl][data-fmt=docx]'))
    pg.fill('[data-cv=summary]','Resumo editado: coração, ação, é, ñ. Sem exageros.');pg.fill('[data-cv="experience.0.bullets"]','Ensinou inglês no ensino secundário.\nTradução técnica EN-PT.');pg.fill('[data-cv=headline]','Professor de inglês')
    pg.check('[data-cvok]');check('confirmar a revisão activa os botões (sem repintar o formulário)',not pg.is_disabled('[data-act=cvdl][data-fmt=pdf]') and pg.input_value('[data-cv=summary]').startswith('Resumo editado'))
    fn,path=download(pg,'[data-act=cvdl][data-fmt=pdf]');t=pdftext(path);r=PdfReader(str(path))
    check('CV em PDF: nome do ficheiro, secções, acentos, edições do utilizador',fn=='CV-Joao-Silva.pdf' and 'EXPERIÊNCIA PROFISSIONAL' in t and 'Resumo editado: coração, ação, é, ñ' in t and 'Joao Silva' in t and 'Escola Y' in t,fn)
    check('CV em PDF: o número removido pelo utilizador já não aparece; sem contacto inventado','120' not in t and '999' not in t and 'invent' not in t.lower())
    subprocess.run(['pdftoppm','-png','-r','60','-f','1','-l','1',str(path),'/tmp/cv-page'])
    pg.select_option('[data-cv=lang]','en');fn,path=download(pg,'[data-act=cvdl][data-fmt=docx]');dt=docxtext(path)
    check('CV em Word: idioma dos títulos EN, texto editado, bullets',fn=='CV-Joao-Silva.docx' and 'PROFESSIONAL EXPERIENCE' in dt and 'Resumo editado' in dt and 'Ensinou inglês no ensino secundário.' in dt,fn)
    fn,pdfpath=download(pg,'[data-act=cvdl][data-fmt=pdf]');check('CV em PDF com títulos em inglês',('PROFESSIONAL EXPERIENCE' in pdftext(pdfpath)) and 'Education'.upper() in pdftext(pdfpath))
    clean(pg,'kit, CV e descargas')
    for fn2,lab in [(pdfpath,'PDF gerado'),(path,'DOCX gerado')]:
        p2=ctx('10.1.6.3');p2.goto(BASE+'/');p2.click('[data-act=start]');p2.set_input_files('#f',str(fn2))
        try:until(p2,"document.querySelector('#cv').value.includes('Resumo editado')",20000);ok=True
        except Exception:ok=False
        check(f'ida e volta: a app lê o {lab} que ela própria criou',ok)
    # 6d) CV gerado sem código => servidor recusa (nem a interface o pede)
    s,j=post({'task':'cv','cv':CV,'vacancy':VAC},'10.1.6.4');check('API: CV adaptado sem código => 402',s==402 and j['code']=='PREMIUM_REQUIRED')
    # 6e) Demo: CV fictício, PDF marcado como DEMO; interface em inglês
    pg=ctx('10.1.6.5');pg.goto(BASE+'/?lang=en');pg.click('[data-act=demo]');pg.wait_for_selector('.score');pg.click('[data-act=kit]');pg.wait_for_selector('.cvform')
    check('demo EN: CV fictício e etiquetas em inglês','Fictional Candidate' in pg.input_value('[data-cv=name]') and pg.locator('text=Download PDF').count()==1 and pg.locator('[data-act=gencv]').count()==0)
    pg.check('[data-cvok]');fn,path=download(pg,'[data-act=cvdl][data-fmt=pdf]');dt=pdftext(path);check('PDF da demo marcado como DEMO',fn=='CV-Fictional-Candidate.pdf' and 'DEMO' in dt and 'Fictional Candidate' in dt,fn)
    # 6f) Sobre + domínio
    pg=ctx('10.1.6.6');pg.goto(BASE+'/');pg.wait_for_selector('#sobre');a=pg.locator('#sobre a[href^="https://snhamoneque"]')
    check('secção Sobre: ligação ao portfólio abre em novo separador, com noopener',a.count()==1 and a.get_attribute('href')=='https://snhamoneque.vercel.app' and a.get_attribute('target')=='_blank' and 'noopener' in a.get_attribute('rel'))
    pg.evaluate("document.querySelector('.hero').dataset.mark='1';scrollTo(0,0)");pg.click('.hero >> text=Como funciona');pg.wait_for_timeout(400)
    check('"Como funciona?" desce até à secção sem repintar a página',pg.evaluate("document.querySelector('.hero').dataset.mark")=='1' and pg.evaluate("document.querySelector('#como').getBoundingClientRect().top<innerHeight*0.6"))
    pg.evaluate('scrollTo(0,0)');pg.click('footer >> text=Sobre');pg.wait_for_timeout(300)
    check('ligação "Sobre" do rodapé leva à secção',pg.evaluate("(()=>{const r=document.querySelector('#sobre').getBoundingClientRect();return r.top>=-5&&r.top<innerHeight})()"))
    pg.click('[data-act=start]');pg.click('footer >> text=Sobre');pg.wait_for_selector('#sobre');check('"Sobre" funciona a partir de outro passo (volta à landing)',pg.locator('#sobre').count()==1)
    html=urllib.request.urlopen(BASE+'/').read().decode()
    check('metadados com o domínio applypro.vercel.app e sem o antigo',html.count('https://applypro.vercel.app')>=4 and 'aijobcopilot' not in html)
    for f in ['robots.txt','sitemap.xml']:
        x=urllib.request.urlopen(BASE+'/'+f).read().decode();check(f+' usa applypro.vercel.app',('applypro.vercel.app' in x) and 'aijobcopilot' not in x)
    # 6g) Fotografar um CV em papel (OCR real, no navegador, sob a CSP estrita)
    if not HAVE_SCAN:
        print('SKIP fotografias: sem tipo de letra DejaVu/Arial para gerar as imagens de teste')
    else:
        def scan_ctx(ip,w=390,h=844):
            pg=ctx(ip,w,h);pg.reqs=[];pg.on('request',lambda r:pg.reqs.append((r.method,r.url,r.post_data or '')));return pg
        def open_scan(pg,qs=''):
            pg.goto(BASE+'/'+qs);pg.click('[data-act=start]');pg.wait_for_selector('.scan')
        pg=scan_ctx('10.1.7.1');open_scan(pg)
        check('cartão de fotografar: câmara nativa (capture=environment) e galeria (multiple), só imagens',pg.get_attribute('#cam','capture')=='environment' and pg.get_attribute('#cam','accept')=='image/*' and pg.get_attribute('#gal','multiple') is not None and 'Fotografar um CV em papel' in pg.inner_text('.scan'))
        pg.set_input_files('#gal',str(FX/'cv-photo.jpg'));pg.wait_for_selector('.pg')
        check('foto escolhida: miniatura (data: URL) com rodar/remover e botão Ler o texto',pg.locator('.pg').count()==1 and pg.get_attribute('.pg img','src').startswith('data:image/jpeg') and pg.locator('[data-act=scanrot]').count()==1 and pg.locator('[data-act=scanrun]').count()==1)
        n0=len(pg.reqs);pg.click('[data-act=scanrun]');scan_done(pg);txt=pg.input_value('#cv');new=pg.reqs[n0:]
        check('OCR real (JPG): texto com acentos, secções e datas',all(k in txt for k in ['Professor de inglês','EXPERIÊNCIA PROFISSIONAL','Universidade Eduardo Mondlane','2020 a 2024','FORMAÇÃO']),txt[:80].replace('\n',' '))
        check('depois da leitura: aviso de revisão, páginas limpas, contador actualizado',pg.inner_text('.scan .okbox').startswith('Texto adicionado') and pg.locator('.pg').count()==0 and str(len(txt)) in pg.inner_text('#c1'))
        check('privacidade: durante a leitura não houve nenhum pedido POST nem /api/, e todos os pedidos foram à própria origem',all(m=='GET' for m,u,b in new) and not any('/api/' in u for m,u,b in new) and all(u.startswith(BASE) for m,u,b in new),str([u for m,u,b in new if not u.startswith(BASE)]))
        used=[u.replace(BASE,'') for m,u,b in new];check('só foram pedidos ficheiros do leitor local (script, worker, motor, língua pt)','/vendor/tesseract/lang/por.traineddata.gz' in used and any('/core/tesseract-core' in u and u.endswith('lstm.wasm.js') for u in used) and '/vendor/tesseract/worker.min.js' in used and not any('eng.traineddata' in u for u in used),str(used))
        clean(pg,'fotografar CV (JPG)')
        for fn,lab in [('cv-photo.png','PNG'),('cv-photo.webp','WebP')]:
            p2=scan_ctx('10.1.7.2');open_scan(p2);p2.set_input_files('#gal',str(FX/fn));p2.wait_for_selector('.pg');p2.click('[data-act=scanrun]');scan_done(p2)
            check(f'OCR real ({lab})','Professor de inglês' in p2.input_value('#cv') and not p2.evaluate('S.scan.err'))
        # segue para a análise: o servidor recebe só o TEXTO, nunca a imagem
        n1=len(pg.reqs);pg.click('[data-act=next1]');pg.fill('#vac',VAC);pg.click('[data-act=next2]');pg.wait_for_selector('.score',timeout=15000)
        post=[b for m,u,b in pg.reqs[n1:] if m=='POST' and '/api/analyze' in u]
        check('análise do CV lido: o pedido leva o texto OCR e nenhuma imagem',len(post)==1 and 'EXPERIÊNCIA PROFISSIONAL' in json.loads(post[0])['cv'] and 'base64' not in post[0] and '/9j/' not in post[0] and len(post[0])<20000,str(len(post[0])) if post else 'sem pedido')
        # rodar
        p3=scan_ctx('10.1.7.3');open_scan(p3);p3.set_input_files('#gal',str(FX/'cv-photo-rot90.jpg'));p3.wait_for_selector('.pg');t0=p3.get_attribute('.pg img','src')
        p3.click('[data-act=scanrot]');t1=p3.get_attribute('.pg img','src');check('rodar: a miniatura muda',t0!=t1)
        p3.click('[data-act=scanrun]');scan_done(p3);check('foto na horizontal: depois de rodar 90° a leitura funciona','EXPERIÊNCIA PROFISSIONAL' in p3.input_value('#cv'),p3.input_value('#cv')[:60].replace('\n',' '))
        # duas páginas (câmara uma a uma + galeria), remover, ordem
        p4=scan_ctx('10.1.7.4');open_scan(p4);p4.set_input_files('#cam',str(FX/'cv-photo.jpg'));p4.wait_for_selector('.pg');p4.set_input_files('#cam',str(FX/'cv-photo-p2.jpg'));until(p4,"S.scan.pages.length==2")
        check('câmara: cada foto acrescenta uma página (campo limpo para poder repetir a mesma foto)',p4.locator('.pg').count()==2 and p4.input_value('#cam')=='')
        p4.click('.pg:nth-child(2) [data-act=scandel]');check('remover: fica 1 página',p4.locator('.pg').count()==1);p4.set_input_files('#gal',str(FX/'cv-photo-p2.jpg'));until(p4,"S.scan.pages.length==2")
        p4.fill('#cv','Notas anteriores do utilizador');p4.click('[data-act=scanrun]');scan_done(p4);t4=p4.input_value('#cv')
        check('2 páginas: junta ao texto que já existia, pela ordem das páginas',t4.startswith('Notas anteriores do utilizador\n\n') and 0<t4.index('EXPERIÊNCIA PROFISSIONAL')<t4.index('CERTIFICADOS') and 'Centro de Línguas' in t4)
        # inglês
        p5=scan_ctx('10.1.7.5');open_scan(p5,'?lang=en');p5.set_input_files('#gal',str(FX/'cv-photo.jpg'));p5.wait_for_selector('.pg')
        check('EN: cartão traduzido e língua por defeito = English',p5.inner_text('.scan h3')=='Photograph a paper CV' and p5.input_value('[data-scanlang]')=='eng')
        n5=len(p5.reqs);p5.click('[data-act=scanrun]');scan_done(p5)
        check('EN: usa os dados de língua English (e não os de português) e conclui',any('eng.traineddata.gz' in u for m,u,b in p5.reqs[n5:]) and not any('por.traineddata' in u for m,u,b in p5.reqs[n5:]) and not p5.evaluate('S.scan.err') and len(p5.input_value('#cv'))>100)
        # erros
        p6=scan_ctx('10.1.7.6');open_scan(p6);p6.set_input_files('#gal',str(FX/'blank.jpg'));p6.wait_for_selector('.pg');p6.click('[data-act=scanrun]');scan_done(p6)
        check('foto sem texto: erro claro, CV intacto, página mantida para tentar de novo',p6.inner_text('.scan .err').startswith('Não foi encontrado texto suficiente') and p6.input_value('#cv')=='' and p6.locator('.pg').count()==1)
        p6.set_input_files('#gal',str(FX/'fake.jpg'));p6.wait_for_selector('.scan .err:has-text("Formato")');check('ficheiro .jpg que não é imagem (conteúdo PDF) é recusado pelos bytes',p6.locator('.pg').count()==1)
        p6.set_input_files('#gal',str(FX/'big.jpg'));p6.wait_for_selector('.scan .err:has-text("demasiado grande")');check('imagem > 15 MB é recusada',p6.locator('.pg').count()==1)
        p6.set_input_files('#gal',[str(FX/'cv-photo.jpg')]*5);until(p6,"S.scan.pages.length==4");check('máximo de 4 páginas (aviso e botões de adicionar desactivados)',p6.inner_text('.scan .err').startswith('Máximo de 4') and p6.is_disabled('#cam') and p6.is_disabled('#gal'))
        check('mobile 390px com 4 miniaturas: sem overflow horizontal',not p6.evaluate("document.documentElement.scrollWidth>document.documentElement.clientWidth"))
        p7=scan_ctx('10.1.7.7',1280,900);open_scan(p7);check('desktop: cartão de fotografar visível sem overflow',p7.locator('.scan').is_visible() and not p7.evaluate("document.documentElement.scrollWidth>document.documentElement.clientWidth"))
        priv=urllib.request.urlopen(BASE+'/privacy.html').read().decode();check('política de privacidade menciona a leitura de fotografias no dispositivo',('lidas (OCR) no teu dispositivo' in priv) and ('read (OCR) on your device' in priv))
    # 7) páginas legais
    pg=ctx('10.1.4.1');pg.goto(BASE+'/?lang=en');pg.wait_for_selector('.hero');pg.click('footer >> text=Privacy');pg.wait_for_selector('.legal')
    vis=pg.evaluate("[...document.querySelectorAll('.legal h1')].filter(h=>h.offsetParent).map(h=>h.innerText)")
    check('ligação do rodapé EN abre a privacidade em inglês (PT escondido)',vis==['Privacy Policy'],str(vis))
    pg.click('[data-lang=pt]');check('páginas legais: botão PT',pg.evaluate("[...document.querySelectorAll('.legal h1')].filter(h=>h.offsetParent).map(h=>h.innerText)")==['Política de Privacidade'])
    pg.goto(BASE+'/terms.html');check('terms.html abre (PT guardado)',pg.title().startswith('Termos'));clean(pg,'legal')
    pg=ctx('10.1.4.2');pg.goto(BASE+'/');pg.wait_for_selector('.hero');href=pg.get_attribute('#sup','href');check('suporte WhatsApp tem ligação válida',href.startswith('https://wa.me/258'))
    # 8) mobile sem overflow
    for url in ['/','/?lang=en','/privacy.html','/terms.html']:
        pg=ctx('10.1.5.1',390,844);pg.goto(BASE+url);pg.wait_for_timeout(400);check(f'mobile 390px sem overflow horizontal {url}',not pg.evaluate("document.documentElement.scrollWidth>document.documentElement.clientWidth"))
    b.close()
srv.terminate();time.sleep(.3);log=LOG.read_text()
check('registos: etiqueta do cliente aparece, os códigos nunca',('"id":"maria"' in log) and not any(x in log for x in ['MARIA-CODE-1234','JOAO-CODE-5678','LEGACYCODE12345','stub-key-not-real']),'')
bad=[n for n,o in R if not o];print(f'\n{len(R)-len(bad)}/{len(R)} verificações passaram');sys.exit(1 if bad else 0)
